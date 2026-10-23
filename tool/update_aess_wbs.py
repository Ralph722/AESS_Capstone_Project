"""Rewrite WBS activity lists to match the real AESS system."""

from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font


path = r"c:\aess_systemproj\assets\Month of September WBS-Development (AESS).xlsx"
wb = openpyxl.load_workbook(path)


def unmerge_in_cols(ws, start_row, end_row, max_col=10):
    """Unmerge any ranges that touch rows start_row..end_row and cols 1..max_col."""
    to_unmerge = []
    for merged in list(ws.merged_cells.ranges):
        if merged.max_row < start_row or merged.min_row > end_row:
            continue
        if merged.max_col < 1 or merged.min_col > max_col:
            continue
        to_unmerge.append(str(merged))
    for rng in to_unmerge:
        ws.unmerge_cells(rng)


def clear_activity_block(ws, start_row, end_row):
    unmerge_in_cols(ws, start_row, end_row, max_col=10)
    for r in range(start_row, end_row + 1):
        for c in range(1, 10):
            cell = ws.cell(r, c)
            # Skip if somehow still a MergedCell
            if getattr(cell, "value", None) is not None or type(cell).__name__ == "Cell":
                try:
                    cell.value = None
                except AttributeError:
                    pass


def write_row(
    ws,
    r,
    task,
    activity,
    assigned=None,
    progress=None,
    start=None,
    end=None,
    bold=False,
):
    ws.cell(r, 1).value = task
    ws.cell(r, 2).value = activity
    if assigned is not None:
        ws.cell(r, 3).value = assigned
    if progress is not None:
        ws.cell(r, 4).value = progress
        ws.cell(r, 4).number_format = "0%"
    if start is not None:
        ws.cell(r, 7).value = start
        ws.cell(r, 7).number_format = "YYYY-MM-DD"
    if end is not None:
        ws.cell(r, 9).value = end
        ws.cell(r, 9).number_format = "YYYY-MM-DD"
    for c in range(1, 10):
        cell = ws.cell(r, c)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if bold:
            cell.font = Font(bold=True)


def d(y, m, day):
    return datetime(y, m, day)


# ============================================================
# AESS - Dev
# ============================================================
ws = wb["AESS - Dev"]
clear_activity_block(ws, 6, 100)
# Clear old signature footer area (known rows from original template)
unmerge_in_cols(ws, 55, 80, max_col=10)
for rr in range(55, 80):
    for c in range(1, 6):
        try:
            ws.cell(rr, c).value = None
        except AttributeError:
            pass

dev_rows = [
    ("1", "Project Initiation", None, None, d(2026, 3, 1), d(2026, 12, 13), True),
    (
        "1.2",
        "Capstone System Development — AESS (Alumni Engagement & Support System)",
        None,
        None,
        None,
        None,
        True,
    ),
    ("1.2.1", "Alumni Portal (Frontend)", None, None, None, None, True),
    (
        "1.2.1.1",
        "Authentication UI (Login / Register / Forgot Password)",
        "RALPH",
        1.0,
        d(2026, 3, 1),
        d(2026, 3, 5),
        False,
    ),
    (
        "1.2.1.2",
        "Alumni eligibility check & registration (roster-linked)",
        "RALPH",
        1.0,
        d(2026, 3, 6),
        d(2026, 3, 10),
        False,
    ),
    (
        "1.2.1.3",
        "Alumni onboarding wizard",
        "RALPH",
        1.0,
        d(2026, 3, 11),
        d(2026, 3, 14),
        False,
    ),
    (
        "1.2.1.4",
        "Alumni Home dashboard",
        "RALPH",
        1.0,
        d(2026, 3, 15),
        d(2026, 3, 18),
        False,
    ),
    (
        "1.2.1.5",
        "Events (browse, RSVP, comments, calendar)",
        "RALPH",
        1.0,
        d(2026, 3, 19),
        d(2026, 3, 24),
        False,
    ),
    (
        "1.2.1.6",
        "Tracer & Surveys (CHED GTS form + custom surveys)",
        "RALPH",
        1.0,
        d(2026, 3, 25),
        d(2026, 4, 5),
        False,
    ),
    (
        "1.2.1.7",
        "Job board (browse, interest, recommendations)",
        "RALPH",
        1.0,
        d(2026, 4, 6),
        d(2026, 4, 12),
        False,
    ),
    (
        "1.2.1.8",
        "Community Forum (topics, replies, tags)",
        "RALPH",
        1.0,
        d(2026, 4, 13),
        d(2026, 4, 18),
        False,
    ),
    (
        "1.2.1.9",
        "Messaging (alumni <-> admin, attachments)",
        "RALPH",
        1.0,
        d(2026, 4, 19),
        d(2026, 4, 24),
        False,
    ),
    (
        "1.2.1.10",
        "Announcements (read / targeted audience)",
        "RALPH",
        1.0,
        d(2026, 4, 25),
        d(2026, 4, 28),
        False,
    ),
    (
        "1.2.1.11",
        "Profile, password change & Alumni ID card",
        "RALPH",
        1.0,
        d(2026, 4, 29),
        d(2026, 5, 5),
        False,
    ),
    (
        "1.2.1.12",
        "Feedback submission & tracking",
        "RALPH",
        1.0,
        d(2026, 5, 6),
        d(2026, 5, 8),
        False,
    ),
    (
        "1.2.1.13",
        "AI Assistant (alumni guidance chatbot)",
        "RALPH",
        1.0,
        d(2026, 5, 9),
        d(2026, 5, 12),
        False,
    ),
    (
        "1.2.1.14",
        "Notifications center & preferences (in-app / push)",
        "RALPH",
        0.8,
        d(2026, 5, 13),
        d(2026, 5, 18),
        False,
    ),
    ("1.3", "Admin Portal (Frontend)", None, None, None, None, True),
    (
        "1.3.1",
        "Admin dashboard shell (sidebar, KPIs, activity)",
        "RALPH",
        1.0,
        d(2026, 5, 19),
        d(2026, 5, 26),
        False,
    ),
    (
        "1.3.2",
        "Alumni Directory & member detail (roster)",
        "RALPH",
        1.0,
        d(2026, 5, 27),
        d(2026, 6, 5),
        False,
    ),
    (
        "1.3.3",
        "Surveys (editor, responses, analytics, export)",
        "RALPH",
        1.0,
        d(2026, 6, 6),
        d(2026, 6, 18),
        False,
    ),
    (
        "1.3.4",
        "Reports & Analytics (GTS charts, presets, CHED compare)",
        "RALPH",
        1.0,
        d(2026, 6, 19),
        d(2026, 6, 30),
        False,
    ),
    (
        "1.3.5",
        "Program Insights / DSS (findings, plan, proof, outcomes)",
        "RALPH",
        1.0,
        d(2026, 7, 1),
        d(2026, 7, 20),
        False,
    ),
    (
        "1.3.6",
        "Events management (CRUD, RSVP, comments)",
        "RALPH",
        1.0,
        d(2026, 7, 21),
        d(2026, 7, 28),
        False,
    ),
    (
        "1.3.7",
        "Announcements management (audience, email delivery)",
        "RALPH",
        1.0,
        d(2026, 7, 29),
        d(2026, 8, 4),
        False,
    ),
    (
        "1.3.8",
        "Forum moderation",
        "RALPH",
        1.0,
        d(2026, 8, 5),
        d(2026, 8, 8),
        False,
    ),
    (
        "1.3.9",
        "Admin messaging inbox",
        "RALPH",
        1.0,
        d(2026, 8, 9),
        d(2026, 8, 12),
        False,
    ),
    (
        "1.3.10",
        "Job postings management & career insights",
        "RALPH",
        1.0,
        d(2026, 8, 13),
        d(2026, 8, 18),
        False,
    ),
    (
        "1.3.11",
        "Activity / audit logs viewer",
        "RALPH",
        1.0,
        d(2026, 8, 19),
        d(2026, 8, 21),
        False,
    ),
    (
        "1.3.12",
        "Archives (events / jobs / surveys)",
        "RALPH",
        1.0,
        d(2026, 8, 22),
        d(2026, 8, 24),
        False,
    ),
    (
        "1.3.13",
        "Feedback management & responses",
        "RALPH",
        1.0,
        d(2026, 8, 25),
        d(2026, 8, 27),
        False,
    ),
    (
        "1.3.14",
        "AI Assistant (admin)",
        "RALPH",
        1.0,
        d(2026, 8, 28),
        d(2026, 8, 29),
        False,
    ),
    (
        "1.3.15",
        "CHED GTS PDF / fillable export",
        "RALPH",
        1.0,
        d(2026, 8, 30),
        d(2026, 9, 5),
        False,
    ),
    ("1.4", "Backend & Platform (Supabase)", None, None, None, None, True),
    (
        "1.4.1",
        "Supabase project setup (Auth, Postgres, Storage, Realtime)",
        "RALPH",
        1.0,
        d(2026, 3, 1),
        d(2026, 3, 8),
        False,
    ),
    (
        "1.4.2",
        "Authentication & role-based access (admin / alumni)",
        "RALPH",
        1.0,
        d(2026, 3, 9),
        d(2026, 3, 14),
        False,
    ),
    (
        "1.4.3",
        "Alumni roster & profiles database (RLS)",
        "RALPH",
        1.0,
        d(2026, 3, 15),
        d(2026, 3, 22),
        False,
    ),
    (
        "1.4.4",
        "Surveys & CHED GTS data model / evaluation",
        "RALPH",
        1.0,
        d(2026, 3, 23),
        d(2026, 4, 5),
        False,
    ),
    (
        "1.4.5",
        "Events, Jobs, Forum, Messaging backends",
        "RALPH",
        1.0,
        d(2026, 4, 6),
        d(2026, 4, 25),
        False,
    ),
    (
        "1.4.6",
        "Announcements + email delivery services",
        "RALPH",
        1.0,
        d(2026, 4, 26),
        d(2026, 5, 2),
        False,
    ),
    (
        "1.4.7",
        "Notifications (in-app, Web Push / FCM)",
        "RALPH",
        0.8,
        d(2026, 5, 3),
        d(2026, 5, 12),
        False,
    ),
    (
        "1.4.8",
        "Audit logging service",
        "RALPH",
        1.0,
        d(2026, 5, 13),
        d(2026, 5, 16),
        False,
    ),
    (
        "1.4.9",
        "Edge Functions (contact verify, exports, DSS helpers)",
        "RALPH",
        1.0,
        d(2026, 5, 17),
        d(2026, 5, 28),
        False,
    ),
    (
        "1.4.10",
        "Program Insights (DSS) server evaluation & rules",
        "RALPH",
        1.0,
        d(2026, 6, 1),
        d(2026, 6, 20),
        False,
    ),
    (
        "1.4.11",
        "Reports / analytics query layer",
        "RALPH",
        1.0,
        d(2026, 6, 21),
        d(2026, 7, 5),
        False,
    ),
    ("1.5", "Cross-cutting Delivery", None, None, None, None, True),
    (
        "1.5.1",
        "Admin UI/UX redesign (shared design system)",
        "RALPH",
        1.0,
        d(2026, 9, 6),
        d(2026, 10, 19),
        False,
    ),
    (
        "1.5.2",
        "Web hosting & release builds (Chrome / Hosting)",
        "RALPH",
        0.5,
        d(2026, 10, 20),
        d(2026, 11, 15),
        False,
    ),
    (
        "1.5.3",
        "System documentation & user/admin guides",
        "RALPH",
        0.7,
        d(2026, 10, 1),
        d(2026, 11, 30),
        False,
    ),
    ("1.6", "Testing and Quality Assurance", None, None, None, None, True),
    (
        "1.6.1",
        "Bugs fixes & optimization",
        "RALPH",
        0.7,
        d(2026, 8, 1),
        d(2026, 11, 30),
        False,
    ),
    (
        "1.6.2",
        "Automated tests (unit / widget / regression SQL)",
        "RALPH",
        0.8,
        d(2026, 8, 1),
        d(2026, 11, 15),
        False,
    ),
    ("1.7", "Project Implementation", None, None, None, None, True),
    (
        "1.7.1",
        "App / web deployment",
        None,
        0.3,
        d(2026, 12, 1),
        d(2026, 12, 5),
        False,
    ),
    (
        "1.7.2",
        "Post-launch monitoring",
        None,
        0.0,
        d(2026, 12, 6),
        d(2026, 12, 10),
        False,
    ),
    (
        "1.7.3",
        "User guide / manual turn-over",
        None,
        0.2,
        d(2026, 12, 8),
        d(2026, 12, 12),
        False,
    ),
    (
        "1.7.4",
        "Project closure",
        None,
        0.0,
        d(2026, 12, 13),
        d(2026, 12, 13),
        False,
    ),
    (
        "1.7.5",
        "Project acceptance",
        None,
        0.0,
        d(2026, 12, 13),
        d(2026, 12, 13),
        False,
    ),
]

first_leaf = last_leaf = None
r = 6
for task, activity, assigned, progress, start, end, bold in dev_rows:
    write_row(ws, r, task, activity, assigned, progress, start, end, bold)
    if progress is not None:
        if first_leaf is None:
            first_leaf = r
        last_leaf = r
    r += 1

last_content_row = r - 1
ws["B1"] = "AESS System Development"
ws["C2"] = (
    f'=(COUNTIF(D{first_leaf}:D{last_leaf},1)'
    f'+COUNTIF(D{first_leaf}:D{last_leaf},"100%"))'
    f"/COUNTA(D{first_leaf}:D{last_leaf})"
)

footer = last_content_row + 3
ws.cell(footer, 2).value = "VALIDATED BY:"
ws.cell(footer, 4).value = "NOTED BY:"
ws.cell(footer + 2, 2).value = "LUTHERLY Q. BONGCAWEL, MIT"
ws.cell(footer + 2, 4).value = "JUBILEE S. DAGA-ANG, MIT"
ws.cell(footer + 3, 2).value = "Capstone Adviser"
ws.cell(footer + 3, 4).value = "IT Program Chairperson"
print("Dev leaf rows", first_leaf, "-", last_leaf, "; last content", last_content_row)

# ============================================================
# Testing
# ============================================================
ws = wb["Testing"]
clear_activity_block(ws, 6, 100)
unmerge_in_cols(ws, 50, 80, max_col=10)
for rr in range(50, 80):
    for c in range(1, 6):
        try:
            ws.cell(rr, c).value = None
        except AttributeError:
            pass

test_rows = [
    ("1", "Project Initiation", None, None, True),
    ("2", "Testing and Quality Assurance — AESS", None, None, True),
    ("2.1", "Alumni Portal Testing", None, None, True),
    ("2.1.1", "Authentication (Login / Register / Forgot Password) testing", "RALPH", 1.0, False),
    ("2.1.2", "Alumni eligibility & registration testing", "RALPH", 1.0, False),
    ("2.1.3", "Onboarding wizard testing", "RALPH", 1.0, False),
    ("2.1.4", "Home dashboard testing", "RALPH", 1.0, False),
    ("2.1.5", "Events (RSVP / comments) testing", "RALPH", 1.0, False),
    ("2.1.6", "Tracer & Surveys (CHED GTS) testing", "RALPH", 1.0, False),
    ("2.1.7", "Job board testing", "RALPH", 1.0, False),
    ("2.1.8", "Community Forum testing", "RALPH", 1.0, False),
    ("2.1.9", "Messaging testing", "RALPH", 1.0, False),
    ("2.1.10", "Announcements testing", "RALPH", 1.0, False),
    ("2.1.11", "Profile & Alumni ID card testing", "RALPH", 1.0, False),
    ("2.1.12", "Feedback & AI Assistant testing", "RALPH", 0.8, False),
    ("2.1.13", "Notifications testing", "RALPH", 0.7, False),
    ("2.2", "Admin Portal Testing", None, None, True),
    ("2.2.1", "Admin dashboard & navigation testing", "RALPH", 1.0, False),
    ("2.2.2", "Alumni Directory / member detail testing", "RALPH", 1.0, False),
    ("2.2.3", "Surveys editor / responses / export testing", "RALPH", 1.0, False),
    ("2.2.4", "Reports & Analytics testing", "RALPH", 1.0, False),
    ("2.2.5", "Program Insights (DSS) workflow testing", "RALPH", 1.0, False),
    ("2.2.6", "Events / Announcements / Forum moderation testing", "RALPH", 1.0, False),
    ("2.2.7", "Jobs management testing", "RALPH", 1.0, False),
    ("2.2.8", "Messages / Feedback / Archives testing", "RALPH", 1.0, False),
    ("2.2.9", "Activity logs & permissions testing", "RALPH", 1.0, False),
    ("2.2.10", "CHED GTS PDF export testing", "RALPH", 1.0, False),
    ("2.3", "Backend & Platform Testing", None, None, True),
    ("2.3.1", "Supabase Auth / RLS testing", "RALPH", 1.0, False),
    ("2.3.2", "Roster & profiles database testing", "RALPH", 1.0, False),
    ("2.3.3", "Surveys / CHED GTS backend testing", "RALPH", 1.0, False),
    ("2.3.4", "Events / Jobs / Forum / Messaging backend testing", "RALPH", 1.0, False),
    ("2.3.5", "Announcements & email delivery testing", "RALPH", 0.8, False),
    ("2.3.6", "Notifications (push) testing", "RALPH", 0.7, False),
    ("2.3.7", "Audit logging testing", "RALPH", 1.0, False),
    ("2.3.8", "Edge Functions testing", "RALPH", 0.8, False),
    ("2.3.9", "DSS evaluation / rules testing", "RALPH", 1.0, False),
    ("2.4", "Bugs Fixes & Optimization", "RALPH", 0.7, False),
    ("2.5", "Project Implementation", None, None, True),
    ("2.5.1", "App / web deployment", None, 0.3, False),
    ("2.5.2", "Post-launch monitoring", None, 0.0, False),
    ("2.5.3", "User guide / manual turn-over", None, 0.2, False),
    ("2.5.4", "Project closure", None, 0.0, False),
    ("2.5.5", "Project acceptance", None, 0.0, False),
]

r = 6
for task, activity, assigned, progress, bold in test_rows:
    write_row(ws, r, task, activity, assigned, progress, None, None, bold)
    r += 1

ws.cell(r + 2, 2).value = "VALIDATED BY:"
ws.cell(r + 4, 2).value = "LUTHERLY Q. BONGCAWEL, MIT"
ws.cell(r + 5, 2).value = "Capstone Adviser"
print("Testing last content row", r - 1)

# ============================================================
# Database Design — AESS areas
# ============================================================
ws = wb["Database Design"]
# Unmerge all merges on this sheet before rewrite
for merged in list(ws.merged_cells.ranges):
    ws.unmerge_cells(str(merged))
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5):
    for cell in row:
        try:
            cell.value = None
        except AttributeError:
            pass

ws["A2"] = "AESS DATABASE AREAS"
ws["B3"] = "TABLE / AREA"
ws["C3"] = "DOMAIN"
ws["D3"] = "CONSTRAINTS / NOTES"
ws["E3"] = "WBS"

db_areas = [
    ("users / auth", "Core", "Supabase Auth + users role", "1.4.2"),
    ("alumni_roster / profiles", "Core", "Official graduates, eligibility", "1.4.3"),
    ("surveys / responses / CHED GTS", "Core", "Tracer waves & answers", "1.4.4"),
    ("events / rsvp / comments", "Community", "Campus events", "1.4.5"),
    ("jobs / job_interest", "Careers", "Job board", "1.4.5"),
    ("forum_topics / replies / tags", "Community", "Moderated discussion", "1.4.5"),
    ("messages / attachments", "Community", "Admin-alumni chat", "1.4.5"),
    ("announcements / deliveries", "Community", "Email + in-app", "1.4.6"),
    ("notifications", "Platform", "In-app / push tokens", "1.4.7"),
    ("audit_logs", "Platform", "Admin accountability", "1.4.8"),
    ("dss_flags / rules / interventions", "Insights", "Program Insights DSS", "1.4.10"),
    ("feedback", "System", "Alumni feedback tickets", "1.3.13"),
    ("archives", "System", "Archived content refs", "1.3.12"),
    ("storage buckets", "Platform", "Images / PDFs / exports", "1.4.1"),
]
for i, (b, c, dnote, e) in enumerate(db_areas, start=4):
    ws.cell(i, 2).value = b
    ws.cell(i, 3).value = c
    ws.cell(i, 4).value = dnote
    ws.cell(i, 5).value = e

# ============================================================
# UAT — retitle for AESS
# ============================================================
ws = wb["UAT"]
for row in ws.iter_rows(min_row=1, max_row=25, max_col=8):
    for cell in row:
        cell.value = None
ws["A1"] = "AESS — User Acceptance Testing log"
ws["A2"] = (
    "Record defects and change requests found during UAT. "
    "Cleared unrelated project (MABISA) placeholders."
)
ws["A5"] = "AESS UAT"
ws["A6"] = "NO."
ws["B6"] = "USER"
ws["C6"] = "MODULE"
ws["D6"] = "TYPE"
ws["E6"] = "ISSUE"
for i in range(1, 8):
    ws.cell(6 + i, 1).value = float(i)
    ws.cell(6 + i, 4).value = "Bug" if i == 1 else "New"
ws.cell(7, 3).value = "Module > Screen > Example"
ws["A20"] = "MIGRATION / DATA NOTES"
ws["A21"] = "NO."
ws["B21"] = "USER"
ws["C21"] = "MODULE"
ws["D21"] = "TYPE"
ws["E21"] = "ISSUE"

wb.save(path)
print("Saved", path)
